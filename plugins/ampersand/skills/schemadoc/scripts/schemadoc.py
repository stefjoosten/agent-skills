#!/usr/bin/env python3
"""schemadoc.py — the schema documentation of an Ampersand context, as one xlsx.

Three living registers in a single workbook (canonical path ./docs/schemaDoc.xlsx):

  sheet 1  Concepts  — one row per concept  (26 columns)
  sheet 2  Relations — one row per relation (24 columns)
  sheet 3  Rules     — one row per rule     (15 columns)
  sheet 4  Legend    — the column semantics of all three sheets

Subcommands:
  scan  <schemaDoc.xlsx> <file.adl>...   build the workbook, or update it in place
  check <schemaDoc.xlsx> <file.adl>...   report differences between workbook and scripts
  gen   <schemaDoc.xlsx> <out.adl>       generate CONCEPT/RELATION/RULE declarations
  pop   <schemaDoc.xlsx> <pop.xlsx>...   analyse a population against Concepts+Relations

One workbook may hold several contexts side by side. A scan or check is scoped to
the context of the scripts it reads, so the rows of the other contexts pass
through untouched — which is what lets a schema live in the workbook before any
script for it exists.

Migration from the former separate matrices: `scan --seed docs/ConceptMatrix.xlsx
--seed docs/RelationMatrix.xlsx ...` reads the manual columns out of the old
workbooks (their sheets are already named Concepts / Relations) and folds them
into the new workbook; rows already present in schemaDoc.xlsx win over seeds.

The parser is text-based and approximates the Ampersand grammar. It is meant for
drafting and consistency checking; when precision matters, verify against the
compiler (ampersand documentation, or the generated concepts.json /
relations.json / rules.json of a prototype). All file I/O is UTF-8.
"""

import argparse
import re
import sys
from pathlib import Path

# Keywords that start a new statement (used to delimit parse windows).
STMT_KEYWORDS = (
    "RELATION", "RULE", "ENFORCE", "INTERFACE", "API", "PURPOSE", "PATTERN",
    "ENDPATTERN", "POPULATION", "CONCEPT", "VIEW", "IDENT", "ROLE", "CLASSIFY",
    "REPRESENT", "MEANING", "MESSAGE", "VIOLATION", "ENDCONTEXT", "INCLUDE",
    "META", "CONTEXT",
)

# ------------------------------------------------------------- sheet layouts

CON_HEADERS = [
    "Name", "Context", "Origin", "Purpose", "Definition", "TIPI",
    "Identity", "Notational conventions", "Example atoms",
    "Generalisations (ISA)", "Specialisations", "Representation",
    "Views", "IDENT statements",
    "Patterns (declared in)",
    "Relations (as source)", "Relations (as target)",
    "Rules (adding atoms)", "Rules (removing atoms)",
    "Interfaces (showing)", "Interfaces (creating atoms)", "Interfaces (deleting atoms)",
    "Roles (create)", "Roles (read)", "Roles (delete)",
    "Resource(s)",
]
C_NAME, C_CONTEXT = 0, 1
C_ORIGIN, C_PURPOSE, C_DEF, C_TIPI = 2, 3, 4, 5
C_IDENTITY, C_NOTATION, C_EXAMPLES, C_ISA, C_RESOURCES = 6, 7, 8, 9, 25
CON_WIDTHS = [20, 14, 20, 40, 46, 24, 30, 26, 28, 18, 18, 14, 20, 18, 18,
              32, 32, 24, 24, 26, 24, 24, 18, 18, 18, 30]

REL_HEADERS = [
    "Name", "Source", "Target", "Context", "Origin", "Purpose", "Meaning",
    "Multiplicity restrictions", "Example pairs",
    "Patterns (declared in)",
    "Rules (used in)", "Rules (adding pairs)", "Rules (removing pairs)",
    "Rules (adding atoms)", "Rules (removing atoms)",
    "Interfaces (mentioning)", "Interfaces (updating)",
    "Interfaces (creating pairs)", "Interfaces (removing pairs)",
    "Roles (create)", "Roles (read)", "Roles (update)", "Roles (delete)",
    "Resource(s)",
]
R_NAME, R_SOURCE, R_TARGET, R_CONTEXT = 0, 1, 2, 3
R_ORIGIN, R_PURPOSE, R_MEANING, R_MULT, R_EXAMPLES, R_RESOURCES = 4, 5, 6, 7, 8, 23
REL_WIDTHS = [22, 14, 14, 14, 20, 40, 40, 12, 30, 18,
              30, 26, 26, 24, 24, 26, 22, 22, 22, 18, 18, 18, 18, 30]

RUL_HEADERS = [
    "Name", "Context", "Origin", "Purpose", "Meaning",
    "Formal term", "Enforcement", "Maintained by",
    "Message", "Violation text", "Example violation",
    "Patterns (declared in)",
    "Relations (used)", "Relations (written)", "Concepts (atoms created/removed)",
]
U_NAME, U_CONTEXT, U_ORIGIN, U_PURPOSE, U_MEANING = 0, 1, 2, 3, 4
U_TERM, U_ENFORCE, U_MAINT, U_MESSAGE, U_VIOLATION, U_EXAMPLE = 5, 6, 7, 8, 9, 10
RUL_WIDTHS = [26, 14, 20, 40, 40, 44, 16, 18, 30, 44, 30, 18, 30, 26, 26]

SHEETS = {
    "Concepts": dict(
        headers=CON_HEADERS, widths=CON_WIDTHS, freeze="C2",
        key=[C_NAME, C_CONTEXT], context_col=C_CONTEXT,
        derived=[0, 1, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24],
        manual=[C_PURPOSE, C_DEF, C_TIPI, C_IDENTITY, C_NOTATION, C_EXAMPLES],
        keep=[C_ORIGIN, C_RESOURCES],
        textcmp=[],
    ),
    "Relations": dict(
        headers=REL_HEADERS, widths=REL_WIDTHS, freeze="E2",
        key=[R_NAME, R_SOURCE, R_TARGET, R_CONTEXT], context_col=R_CONTEXT,
        derived=[0, 1, 2, 3, 7, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22],
        manual=[R_PURPOSE, R_MEANING, R_EXAMPLES],
        keep=[R_ORIGIN, R_RESOURCES],
        textcmp=[],
    ),
    "Rules": dict(
        headers=RUL_HEADERS, widths=RUL_WIDTHS, freeze="C2",
        key=[U_NAME, U_CONTEXT], context_col=U_CONTEXT,
        derived=[0, 1, 5, 6, 7, 8, 9, 11, 12, 13, 14],
        manual=[U_PURPOSE, U_MEANING, U_EXAMPLE],
        keep=[U_ORIGIN],
        textcmp=[U_TERM, U_MESSAGE, U_VIOLATION],
    ),
}

LEGEND = [
    ("CONCEPTS", ""),
    ("Name / Context", "The concept name identifies the concept within its context (upper case start, singular). The workbook may hold several contexts side by side; then Name and Context together identify the row, and a scan of one context leaves the rows of the others untouched."),
    ("Origin", "The file (relative path) that inspired this concept; scan fills the declaring script when the cell is empty and leaves a filled cell alone."),
    ("Purpose", "Why must atoms of this concept be registered (the PURPOSE CONCEPT statement)?"),
    ("Definition", "The criterion that decides whether something is an atom of this concept (the CONCEPT statement). Good as long as it works."),
    ("TIPI", "Manual verdict on the four test questions — Telbaar, Instantieerbaar, Permanent, Identificeerbaar — with notes on the choices made (esp. permanence within the context)."),
    ("Identity", "The criterion that defines the equality of two atoms of this concept, in prose. Manual; an IDENT statement is its formalisation."),
    ("Notational conventions", "Naming conventions for the atoms (e.g. \"first name, optional hyphen and number\"). Manual."),
    ("Example atoms", "Up to three atoms \"a\", \"b\", \"c\" — real if possible, fictitious is fine."),
    ("Generalisations / Specialisations", "Concepts this one specialises (CLASSIFY this ISA that, or CLASSIFY this IS ...) and the derived converse."),
    ("Representation", "The REPRESENT ttype; empty means the default: ALPHANUMERIC, or OBJECT when an interface boxes the concept (atoms are identities, not values)."),
    ("Views / IDENT statements", "VIEW statements on the concept (DEFAULT marked) and IDENT statements (identifying attribute combinations, the formalisation of the Identity column)."),
    ("Relations (as source/target)", "Declared relations typed with this concept. Both cells empty = the concept is used by nothing — unfinished thinking or dead weight."),
    ("Rules (adding/removing atoms)", "ExecEngine violations: InsAtom and NewStruct create an atom of the concept; DelAtom and MrgAtoms remove one. String-based, so verify by hand."),
    ("Interfaces (showing/creating/deleting)", "From the crud annotation of fields whose target is this concept (upper case = allowed): R shows atoms, C creates an atom, D deletes an atom and its pairs. U is about pairs and lives in the Relations sheet."),
    ("Roles (create/read/delete)", "Union of the interface rights per role, taken from the FOR clauses. An interface without FOR grants nothing, not even to the anonymous session."),
    ("", ""),
    ("RELATIONS", ""),
    ("Name / Source / Target / Context", "The signature Name[Source*Target] identifies the relation within its context; the name alone is not unique. With several contexts in the workbook the context is part of the identity too."),
    ("Origin", "The file (relative path) that inspired this relation; scan fills the declaring script when the cell is empty and leaves a filled cell alone."),
    ("Purpose", "Why does this relation exist (the PURPOSE RELATION statement)?"),
    ("Meaning", "The criterion that decides whether a pair belongs in the relation (the MEANING statement)."),
    ("Multiplicity restrictions", "Comma-separated list: UNI, TOT, INJ, SUR; for endorelations also SYM, ASY, TRN, RFX, IRF, PROP; shorthands MAP, BIJ. The compiler turns each of these into a property rule, checked as an invariant — the Rules sheet does not repeat them."),
    ("Example pairs", "Up to three pairs (\"a\",\"b\"), (\"c\",\"d\") — real if possible, fictitious is fine."),
    ("Rules (used in)", "Rules and ENFORCE statements whose term mentions the relation. The names match the Name column of the Rules sheet."),
    ("Rules (adding/removing pairs)", "ENFORCE := adds and removes; >: adds; :< removes. ExecEngine violations count via InsPair (and the pair inserts of NewStruct); DelPair removes. String-based, so verify by hand."),
    ("Rules (adding/removing atoms)", "ExecEngine InsAtom and the atom created by NewStruct add an atom; DelAtom and MrgAtoms remove one. Attributed to every relation whose source or target is the concept involved."),
    ("Interfaces (mentioning)", "Interfaces and APIs whose field terms mention the relation."),
    ("Interfaces (updating/creating/removing)", "From the crud annotation of the field (upper case = allowed): U = insert and remove pairs; C = create a target atom (and thereby pairs); D = delete a target atom (and thereby its pairs). Only a field whose term is a single, possibly flipped, relation is truly editable."),
    ("Roles (create/read/update/delete)", "Union of the interface rights per role, taken from the FOR clauses."),
    ("", ""),
    ("RULES", ""),
    ("Name / Context", "The rule label; unique within its context (the compiler requires that). An ENFORCE statement has no label and appears as 'ENFORCE <relation> <operator>'. An unlabelled RULE gets a synthetic name from its term — give rules a label to keep the register stable."),
    ("Origin", "The file (relative path) that declares this rule; scan fills it when the cell is empty and leaves a filled cell alone."),
    ("Purpose", "Why must this rule hold (the PURPOSE RULE statement)? The reason of existence: which agreement, law, or design decision demands it."),
    ("Meaning", "The constraint in natural language (the MEANING statement), stated so a stakeholder can judge any concrete case against it."),
    ("Formal term", "The relation-algebra term that must hold; for an ENFORCE row: '<relation> <operator> <term>'. A rule states WHAT must be true, never HOW it is enforced."),
    ("Enforcement", "How the system responds to a violation. invariant = the transaction that causes a violation is blocked (the default). process = a user role gets a persistent signal and must resolve it (ROLE <role> MAINTAINS). automated = the ExecEngine restores the rule via the {EX} actions in its VIOLATION (ROLE ExecEngine MAINTAINS). ENFORCE := / >: / :< = declarative automation: the relation is kept equal to / a superset of / a subset of the term."),
    ("Maintained by", "The role(s) from ROLE ... MAINTAINS. Empty for an invariant: the system itself blocks. ExecEngine for automated rules and ENFORCE."),
    ("Message", "The MESSAGE text(s): the feedback a user sees when the rule is violated."),
    ("Violation text", "The VIOLATION segments (TXT / SRC / TGT), which compose a specific message per violating pair — or, for automated rules, the {EX} actions (InsPair, DelPair, NewStruct, InsAtom, DelAtom, MrgAtoms) the ExecEngine runs per violation."),
    ("Example violation", "Manual: a concrete scenario that violates the rule (and, for process/automated rules, how it gets resolved). The counterpart of Example pairs/atoms: it makes the rule testable in conversation with a stakeholder."),
    ("Relations (used)", "Relations mentioned in the formal term — the rows of the Relations sheet this rule reads."),
    ("Relations (written)", "Relations the enforcement writes: the ENFORCE target, and the relations in InsPair/DelPair/NewStruct actions."),
    ("Concepts (atoms created/removed)", "Concepts whose atoms the enforcement creates (InsAtom, NewStruct) or removes (DelAtom, MrgAtoms)."),
    ("", ""),
    ("CONVENTIONS", ""),
    ("Empty cell vs '-'", "In the manual columns an empty cell means \"not yet analysed\" and a single '-' means \"verified empty\". In the script-derived columns an empty cell means the scan found no occurrence. Rows are kept when their subject disappears from the scripts; 'check' reports them."),
    ("Default crud", "A field without annotation gets the project default (assumed cRud = read only; override with --default-crud)."),
    ("Property rules", "Multiplicity annotations ([UNI], [TOT], ... on a relation) compile into invariant rules of their own. They live in the Multiplicity column of the Relations sheet and are deliberately not repeated as rows in the Rules sheet."),
]


# ---------------------------------------------------------------- text handling

def strip_comments(text: str) -> str:
    """Remove -- line comments and {- -} block comments (nested), preserving
    strings and {+ +} markup blocks. Removed characters become spaces so that
    positions and line structure survive."""
    out = list(text)
    i, n = 0, len(text)
    depth = 0
    in_str = in_markup = in_line = False
    while i < n:
        c = text[i]
        two = text[i:i + 2]
        if in_line:
            if c == "\n":
                in_line = False
            else:
                out[i] = " "
        elif depth > 0:
            if two == "{-":
                depth += 1
                out[i] = out[i + 1] = " "
                i += 1
            elif two == "-}":
                depth -= 1
                out[i] = out[i + 1] = " "
                i += 1
            elif c != "\n":
                out[i] = " "
        elif in_str:
            if c == '"':
                in_str = False
        elif in_markup:
            if two == "+}":
                in_markup = False
                i += 1
        else:
            if two == "{-":
                depth = 1
                out[i] = out[i + 1] = " "
                i += 1
            elif two == "--":
                in_line = True
                out[i] = " "
            elif two == "{+":
                in_markup = True
                i += 1
            elif c == '"':
                in_str = True
        i += 1
    return "".join(out)


def unquote(s: str) -> str:
    s = s.strip()
    return s[1:-1] if len(s) >= 2 and s[0] == '"' and s[-1] == '"' else s


def window(text: str, start: int) -> str:
    """Text from start up to the next statement keyword (exclusive)."""
    m = re.search(r"\b(" + "|".join(STMT_KEYWORDS) + r")\b", text[start:])
    return text[start:start + m.start()] if m else text[start:]


def rule_extent(text: str, start: int) -> str:
    """Text from start through the MEANING/MESSAGE/VIOLATION clauses that belong
    to the same RULE statement (they are statement keywords themselves, so a
    single window stops too early)."""
    pos = start
    while True:
        w = window(text, pos)
        pos += len(w)
        m = re.match(r"(?:MEANING|MESSAGE|VIOLATION)\b", text[pos:pos + 10])
        if not m:
            return text[start:pos]
        pos += m.end()


def violation_text(extent: str) -> str:
    """The content of the VIOLATION ( ... ) clause, brackets balanced and
    quote-aware, whitespace normalised. Empty when there is none."""
    i = extent.find("VIOLATION")
    if i < 0:
        return ""
    j = extent.find("(", i)
    if j < 0:
        return ""
    depth, k, in_str = 0, j, False
    while k < len(extent):
        c = extent[k]
        if in_str:
            if c == '"':
                in_str = False
        elif c == '"':
            in_str = True
        elif c == "(":
            depth += 1
        elif c == ")":
            depth -= 1
            if depth == 0:
                return re.sub(r"\s+", " ", extent[j + 1:k]).strip()
        k += 1
    return re.sub(r"\s+", " ", extent[j + 1:]).strip()


def norm_text(v: str) -> str:
    return re.sub(r"\s+", " ", str(v)).strip()


# ---------------------------------------------------------------- data model

class Con:
    def __init__(self, name):
        self.name = name
        self.definition = ""
        self.purpose = ""
        self.examples = []
        self.isa = set()
        self.spec = set()
        self.represent = ""
        self.views = set()
        self.idents = set()
        self.origin = set()
        self.context = ""
        self.patterns = set()
        self.rels_source = set()
        self.rels_target = set()
        self.rules_addatom = set()
        self.rules_delatom = set()
        self.ifc_show = set()
        self.ifc_create = set()
        self.ifc_delete = set()
        self.roles = {"C": set(), "R": set(), "D": set()}

    @property
    def key(self):
        return (self.name, self.context)


class Rel:
    def __init__(self, name, source, target):
        self.name, self.source, self.target = name, source, target
        self.props = ""
        self.meaning = ""
        self.purpose = ""
        self.examples = []          # list of (a, b)
        self.origin = set()
        self.context = ""
        self.patterns = set()
        self.rules_used = set()
        self.rules_add = set()
        self.rules_del = set()
        self.rules_addatom = set()
        self.rules_delatom = set()
        self.ifc_mention = set()
        self.ifc_update = set()
        self.ifc_create = set()
        self.ifc_remove = set()
        self.roles = {"C": set(), "R": set(), "U": set(), "D": set()}

    @property
    def sig(self):
        return (self.name, self.source, self.target)

    @property
    def key(self):
        return (self.name, self.source, self.target, self.context)


class Rul:
    def __init__(self, name):
        self.name = name
        self.purpose = ""
        self.meaning = ""
        self.term = ""
        self.enforcement = "invariant"
        self.maintainers = set()
        self.messages = set()
        self.violation = ""
        self.origin = set()
        self.context = ""
        self.patterns = set()
        self.rels_used = set()
        self.rels_written = set()
        self.concepts_touched = set()

    @property
    def key(self):
        return (self.name, self.context)


# ---------------------------------------------------------------- parsing

CPT = r"(?:\"[A-Za-z_][\w. ]*\"|[A-Za-z_]\w*)"
CONCEPT_RE = re.compile(r"\bCONCEPT\s+(" + CPT + r")\s+\"((?:[^\"\\]|\\.)*)\"")
CLASSIFY_ISA_RE = re.compile(r"\bCLASSIFY\s+(" + CPT + r")\s+ISA\s+(" + CPT + r")")
CLASSIFY_IS_RE = re.compile(r"\bCLASSIFY\s+(" + CPT + r")\s+IS\s+([^\n]+)")
REPRESENT_RE = re.compile(r"\bREPRESENT\s+([\w\", .]+?)\s+TYPE\s+(\w+)")
VIEW_RE = re.compile(r"\bVIEW\s+(" + CPT + r")\s*(?:LABEL\s+\"[^\"]+\"\s*)?:\s*(" + CPT + r")\s*(DEFAULT)?")
IDENT_RE = re.compile(r"\bIDENT\s+(" + CPT + r")\s*:\s*(" + CPT + r")\s*(\([^)]*\))?")
REL_RE = re.compile(
    r"\bRELATION\s+([A-Za-z_][\w.]*)\s*"
    r"\[\s*\"?([A-Za-z_][\w. ]*?)\"?\s*(?:\*\s*\"?([A-Za-z_][\w. ]*?)\"?\s*)?\]"
    r"(?:\s*\[\s*([A-Z, ]+?)\s*\])?"
)
MEANING_RE = re.compile(
    r"\bMEANING\s+(?:IN\s+\w+\s+)?(?:(?:LATEX|MARKDOWN|HTML|REST)\s+)?"
    r"(?:\{\+(.*?)\+\}|\"((?:[^\"\\]|\\.)*)\")", re.S)
MESSAGE_RE = re.compile(
    r"\bMESSAGE\s+(?:IN\s+\w+\s+)?(?:\{\+(.*?)\+\}|\"((?:[^\"\\]|\\.)*)\")", re.S)
PURPOSE_RE = re.compile(
    r"\bPURPOSE\s+(CONCEPT|RELATION|RULE)\s+(\"[^\"]+\"|[A-Za-z_][\w.]*)"
    r"(?:\s*\[\s*\"?([A-Za-z_][\w. ]*?)\"?\s*(?:\*\s*\"?([A-Za-z_][\w. ]*?)\"?\s*)?\])?"
    r"\s*(?:IN\s+\w+\s*)?(?:(?:LATEX|MARKDOWN|HTML|REST)\s*)?(?:REF\s+\"[^\"]*\"\s*)?"
    r"\{\+(.*?)\+\}", re.S)
POP_RE = re.compile(
    r"\bPOPULATION\s+(" + CPT + r")"
    r"(?:\s*\[\s*\"?([A-Za-z_][\w. ]*?)\"?\s*(?:\*\s*\"?([A-Za-z_][\w. ]*?)\"?\s*)?\])?"
    r"\s+CONTAINS\s*\[(.*?)\]", re.S)
PAIR_RE = re.compile(r"\(\s*\"([^\"]*)\"\s*,\s*\"([^\"]*)\"\s*\)")
ENFORCE_RE = re.compile(r"\bENFORCE\s+([A-Za-z_][\w.]*)(?:\s*\[[^\]]*\])?\s*(:=|>:|:<)")
ROLE_RE = re.compile(r"\bROLE\s+([^\n]+?)\s+MAINTAINS\s+([^\n]+)")
IFC_RE = re.compile(
    r"\b(?:TRANSACTIONAL\s+)?(INTERFACE|API)\s+(\"[^\"]+\"|[A-Za-z_]\w*)"
    r"(?:\s+LABEL\s+\"[^\"]+\")?"
    r"(?:\s+FOR\s+([^:\n]+))?"
    r"\s*:")
PATTERN_RE = re.compile(r"\bPATTERN\s+(\"[^\"]+\"|[A-Za-z_][\w]*)")
CONTEXT_RE = re.compile(r"\bCONTEXT\s+(\"[^\"]+\"|[A-Za-z_][\w]*)")
CRUD_RE = re.compile(r"\b([cC][rR][uU][dD])\b")
TOKEN_RE = re.compile(r"\b(?:[A-Z]\w*\.)?[a-z]\w*\b")
RELTOK_RE = re.compile(r"\b((?:[A-Z]\w*\.)?[a-z]\w*)\s*(~?)")
IV_RE = re.compile(r"\b[IV]\[\s*(" + CPT + r")\s*(?:\*\s*(" + CPT + r")\s*)?\]")
# Een regel heet `RULE naam : term`, of sinds v5.9 ook `RULE naam LABEL "..." : term`.
RULE_LABEL_RE = re.compile(
    r"\s*(?:\"([^\"]+)\"|([A-Za-z_][\w ]*?))\s*(?:LABEL\s+\"[^\"]+\"\s*)?:")


def rel_tokens(term: str, by_name: dict) -> set:
    """Declared relations mentioned in a term. Matches full or unqualified names."""
    found = set()
    for tok in TOKEN_RE.findall(term):
        short = tok.split(".")[-1]
        for key in by_name.get(tok, ()) or by_name.get(short, ()):
            found.add(key)
    return found


def parse_violation_actions(vtext: str):
    """Split a violation text on {EX} and yield (function, argument-string)."""
    for act in vtext.split("{EX}")[1:] if "{EX}" in vtext else []:
        m = re.match(r'[\s"]*([A-Za-z_]\w*)\s*;', act)
        if m:
            yield m.group(1), act


def parse(paths, default_crud="cRud", context_override=None):
    """Parse the scripts into (cons, rels, ruls, context)."""
    texts = {}
    for p in paths:
        texts[p] = strip_comments(Path(p).read_text(encoding="utf-8"))

    rule_purpose = {}   # rule label -> PURPOSE RULE body, toegekend na het lezen van de regels
    cons = {}           # name -> Con
    rels = {}           # (name, src, tgt) -> Rel
    ruls = {}           # name -> Rul
    by_name = {}        # relation name -> [signatures]
    rel_types = {}      # relation name -> set of (src, tgt)
    role_maint = {}     # rule label -> set of roles
    context = context_override or ""

    def con(name, origin=None, pattern=None):
        c = cons.setdefault(name, Con(name))
        if origin:
            c.origin.add(origin)
        if pattern:
            c.patterns.add(pattern)
        return c

    pat_spans_of = {}

    # Pass 1: context, patterns, declarations, ROLE ... MAINTAINS.
    for p, text in texts.items():
        if not context:
            m = CONTEXT_RE.search(text)
            if m:
                context = unquote(m.group(1))
        pat_spans = []
        for m in PATTERN_RE.finditer(text):
            end = text.find("ENDPATTERN", m.end())
            pat_spans.append((m.start(), end if end != -1 else len(text), unquote(m.group(1))))
        pat_spans_of[p] = pat_spans

        def pattern_at(pos):
            for a, b, name in pat_spans:
                if a <= pos < b:
                    return name
            return None

        for m in CONCEPT_RE.finditer(text):
            c = con(unquote(m.group(1)), p, pattern_at(m.start()))
            c.definition = norm_text(m.group(2))

        for m in CLASSIFY_ISA_RE.finditer(text):
            child, parent = unquote(m.group(1)), unquote(m.group(2))
            con(child, p, pattern_at(m.start())).isa.add(parent)
            con(parent, p).spec.add(child)

        for m in CLASSIFY_IS_RE.finditer(text):
            child = unquote(m.group(1))
            parents = re.findall(r"\b[A-Z]\w*\b", m.group(2))
            c = con(child, p, pattern_at(m.start()))
            for parent in parents:
                c.isa.add(parent)
                con(parent, p).spec.add(child)

        for m in REPRESENT_RE.finditer(text):
            for name in m.group(1).split(","):
                name = unquote(name)
                if name:
                    con(name, p, pattern_at(m.start())).represent = m.group(2).strip()

        for m in VIEW_RE.finditer(text):
            label = unquote(m.group(1))
            if m.group(3):
                label += " (DEFAULT)"
            con(unquote(m.group(2)), p).views.add(label)

        for m in IDENT_RE.finditer(text):
            label = unquote(m.group(1))
            if m.group(3):
                label += " " + norm_text(m.group(3))
            con(unquote(m.group(2)), p).idents.add(label)

        for m in REL_RE.finditer(text):
            name, src = m.group(1), m.group(2).strip()
            tgt = (m.group(3) or m.group(2)).strip()
            r = rels.setdefault((name, src, tgt), Rel(name, src, tgt))
            if m.group(4):
                r.props = ", ".join(t.strip() for t in m.group(4).split(","))
            # MEANING is itself a statement keyword, so the window ends just
            # before it; only accept a MEANING that starts where the window ends.
            w = window(text, m.end())
            mm = MEANING_RE.search(text[m.end():m.end() + len(w) + 4000])
            if mm and mm.start() <= len(w) + 1:
                r.meaning = norm_text(mm.group(1) or mm.group(2) or "")
            pat = pattern_at(m.start())
            if pat:
                r.patterns.add(pat)
            r.origin.add(p)
            by_name.setdefault(name, []).append(r.sig)
            sig = f"{name}[{src}*{tgt}]"
            con(src, p, pat).rels_source.add(sig)
            con(tgt, p, pat).rels_target.add(sig)
            rel_types.setdefault(name, set()).add((src, tgt))
            rel_types.setdefault(name.split(".")[-1], set()).add((src, tgt))

        for m in ROLE_RE.finditer(text):
            roles = [unquote(t) for t in m.group(1).split(",") if t.strip()]
            for label in m.group(2).split(","):
                label = unquote(label)
                if label:
                    role_maint.setdefault(label, set()).update(roles)

    by_concept = {}
    for key, r in rels.items():
        by_concept.setdefault(r.source, set()).add(key)
        by_concept.setdefault(r.target, set()).add(key)

    def resolve(name, src=None, tgt=None):
        if src:
            tgt = tgt or src
            return [(name, src, tgt)] if (name, src, tgt) in rels else []
        return by_name.get(name, []) or by_name.get(name.split(".")[-1], [])

    def apply_actions(label, vtext, rul=None):
        """Attribute the {EX} actions of a violation to relations and concepts,
        and (when given) to the rule row itself."""
        for fn, act in parse_violation_actions(vtext):
            arg = re.match(r'[\s"]*\w+\s*;\s*([\w.]+)', act)
            first = arg.group(1) if arg else ""
            # The action fragment starts inside a TXT string, so the quote
            # characters toggle: the even-indexed parts of split('"') are the
            # in-string text. Only relation names THERE are written; a name in
            # a SRC/TGT expression supplies an atom value and is merely read.
            instring = "".join(act.split('"')[0::2])
            if fn in ("InsPair", "SetToday"):
                for key in resolve(first):
                    rels[key].rules_add.add(label)
                    if rul:
                        rul.rels_written.add(f"{key[0]}[{key[1]}*{key[2]}]")
            elif fn == "DelPair":
                for key in resolve(first):
                    rels[key].rules_del.add(label)
                    if rul:
                        rul.rels_written.add(f"{key[0]}[{key[1]}*{key[2]}]")
            elif fn == "NewStruct":
                if first in cons:
                    cons[first].rules_addatom.add(label)
                    if rul:
                        rul.concepts_touched.add(first + " (+)")
                for key in by_concept.get(first, ()):
                    rels[key].rules_addatom.add(label)
                for key in rel_tokens(instring.replace(";", " "), by_name):
                    rels[key].rules_add.add(label)
                    if rul:
                        rul.rels_written.add(f"{key[0]}[{key[1]}*{key[2]}]")
            elif fn == "InsAtom":
                if first in cons:
                    cons[first].rules_addatom.add(label)
                    if rul:
                        rul.concepts_touched.add(first + " (+)")
                for key in by_concept.get(first, ()):
                    rels[key].rules_addatom.add(label)
            elif fn in ("DelAtom", "MrgAtoms"):
                names = {first}
                if fn == "MrgAtoms":
                    names.update(c for c in re.findall(r";([A-Z]\w*);", instring)
                                 if c in cons)
                for name in names:
                    if name in cons:
                        cons[name].rules_delatom.add(label)
                        if rul:
                            rul.concepts_touched.add(name + " (-)")
                    for key in by_concept.get(name, ()):
                        rels[key].rules_delatom.add(label)

    # Pass 2: purposes, populations, rules, enforce, interfaces.
    for p, text in texts.items():
        pat_spans = pat_spans_of[p]

        def pattern_at(pos):
            for a, b, name in pat_spans:
                if a <= pos < b:
                    return name
            return None

        for m in PURPOSE_RE.finditer(text):
            kind, name = m.group(1), unquote(m.group(2))
            body = norm_text(m.group(5))
            if kind == "CONCEPT" and name in cons:
                cons[name].purpose = body
            elif kind == "RELATION":
                for key in resolve(name, m.group(3) and m.group(3).strip(),
                                   m.group(4) and m.group(4).strip()):
                    rels[key].purpose = body
            elif kind == "RULE":
                # De regels zelf worden verderop gelezen; bewaar de purpose zolang.
                rule_purpose[name] = body

        for m in POP_RE.finditer(text):
            name = unquote(m.group(1))
            body = m.group(4)
            pairs = PAIR_RE.findall(body)
            if pairs or name[:1].islower():
                for key in resolve(name, m.group(2) and m.group(2).strip(),
                                   m.group(3) and m.group(3).strip()):
                    if pairs and not rels[key].examples:
                        rels[key].examples = pairs[:3]
                    for concept, idx in ((key[1], 0), (key[2], 1)):
                        c = cons.get(concept)
                        if c is not None and not c.examples:
                            atoms = [pr[idx].strip() for pr in pairs if pr[idx].strip()]
                            if atoms:
                                c.examples = atoms[:3]
                continue
            atoms = re.findall(r"\"([^\"]*)\"", body)
            if atoms and name in cons and not cons[name].examples:
                cons[name].examples = atoms[:3]

        for m in ENFORCE_RE.finditer(text):
            relname, op = m.group(1), m.group(2)
            label = f"ENFORCE {relname} {op}"
            w = window(text, m.end())
            term = norm_text(w)
            u = ruls.setdefault(label, Rul(label))
            u.term = f"{relname} {op} {term}"
            u.enforcement = f"ENFORCE {op}"
            u.maintainers.add("ExecEngine")
            u.origin.add(p)
            pat = pattern_at(m.start())
            if pat:
                u.patterns.add(pat)
            for key in resolve(relname):
                r = rels[key]
                r.rules_used.add(label)
                u.rels_written.add(f"{key[0]}[{key[1]}*{key[2]}]")
                if op in (":=", ">:"):
                    r.rules_add.add(label)
                if op in (":=", ":<"):
                    r.rules_del.add(label)
            for key in rel_tokens(w, by_name):
                rels[key].rules_used.add(label)
                u.rels_used.add(f"{key[0]}[{key[1]}*{key[2]}]")

        for m in re.finditer(r"\bRULE\b", text):
            # `PURPOSE RULE <naam> {+ ... +}` declareert geen regel. Zonder deze test
            # leest de scanner een dubbele punt in de purpose-tekst als regeldefinitie.
            if re.search(r"\bPURPOSE\s*$", text[:m.start()]):
                continue
            head = text[m.end():m.end() + 200]
            lm = RULE_LABEL_RE.match(head)
            label = unquote(lm.group(1) or lm.group(2)) if lm else ""
            w = window(text, m.end())
            term = norm_text(w[lm.end():] if lm else w)
            if not label:
                label = "RULE " + term[:40]
            extent = rule_extent(text, m.end())
            u = ruls.setdefault(label, Rul(label))
            u.term = term
            u.origin.add(p)
            pat = pattern_at(m.start())
            if pat:
                u.patterns.add(pat)
            maint = role_maint.get(label, set())
            if maint:
                u.maintainers.update(maint)
                u.enforcement = ("automated" if any(r.lower() == "execengine" for r in maint)
                                 else "process")
            mm = MEANING_RE.search(extent)
            if mm:
                u.meaning = norm_text(mm.group(1) or mm.group(2) or "")
            for sm in MESSAGE_RE.finditer(extent):
                u.messages.add(norm_text(sm.group(1) or sm.group(2) or ""))
            vtext = violation_text(extent)
            u.violation = vtext
            for key in rel_tokens(w, by_name):
                rels[key].rules_used.add(label)
                u.rels_used.add(f"{key[0]}[{key[1]}*{key[2]}]")
            if vtext:
                apply_actions(label, vtext, u)

        for m in IFC_RE.finditer(text):
            ifc = unquote(m.group(2))
            roles = [t.strip() for t in (m.group(3) or "").split(",") if t.strip()]
            body = window(text, m.end())
            # Fields: split on commas/braces; associate a crud annotation with
            # the relations/concepts in the same field fragment. A field shows
            # the target concept of its term (the source when flipped with ~).
            for frag in re.split(r"[,{}]", body):
                keys = rel_tokens(frag, by_name)
                shown = set()
                for tok, flip in RELTOK_RE.findall(frag):
                    short = tok.split(".")[-1]
                    for src, tgt in rel_types.get(tok, ()) or rel_types.get(short, ()):
                        shown.add(src if flip else tgt)
                for mm in IV_RE.finditer(frag):
                    shown.add(unquote(mm.group(2) or mm.group(1)))
                if not keys and not shown:
                    continue
                cm = CRUD_RE.search(frag)
                crud = cm.group(1) if cm else default_crud
                for key in keys:
                    r = rels[key]
                    r.ifc_mention.add(ifc)
                    if crud[2] == "U":
                        r.ifc_update.add(ifc)
                    if crud[0] == "C":
                        r.ifc_create.add(ifc)
                    if crud[2] == "U" or crud[3] == "D":
                        r.ifc_remove.add(ifc)
                    for letter, allowed in zip("CRUD", crud):
                        if allowed.isupper():
                            r.roles[letter].update(roles)
                for name in shown:
                    if name not in cons:
                        continue
                    c = cons[name]
                    if crud[1] == "R":
                        c.ifc_show.add(ifc)
                        c.roles["R"].update(roles)
                    if crud[0] == "C":
                        c.ifc_create.add(ifc)
                        c.roles["C"].update(roles)
                    if crud[3] == "D":
                        c.ifc_delete.add(ifc)
                        c.roles["D"].update(roles)

    for name, body in rule_purpose.items():
        if name in ruls:
            ruls[name].purpose = body

    # Complete the ISA converse and stamp the context.
    for c in cons.values():
        for parent in c.isa:
            if parent in cons:
                cons[parent].spec.add(c.name)
    for coll in (cons, rels, ruls):
        for x in coll.values():
            x.context = context
    return cons, rels, ruls, context


# ---------------------------------------------------------------- rows

def fmt_set(s):
    return ", ".join(sorted(s))


def fmt_pairs(pairs):
    return ", ".join(f'("{a}","{b}")' for a, b in pairs)


def fmt_atoms(atoms):
    return ", ".join(f'"{a}"' for a in atoms)


def row_of_con(c):
    return [
        c.name, c.context, fmt_set(c.origin), c.purpose, c.definition, "", "", "",
        fmt_atoms(c.examples), fmt_set(c.isa), fmt_set(c.spec), c.represent,
        fmt_set(c.views), fmt_set(c.idents), fmt_set(c.patterns),
        fmt_set(c.rels_source), fmt_set(c.rels_target),
        fmt_set(c.rules_addatom), fmt_set(c.rules_delatom),
        fmt_set(c.ifc_show), fmt_set(c.ifc_create), fmt_set(c.ifc_delete),
        fmt_set(c.roles["C"]), fmt_set(c.roles["R"]), fmt_set(c.roles["D"]),
        "",
    ]


def row_of_rel(r):
    return [
        r.name, r.source, r.target, r.context, fmt_set(r.origin), r.purpose,
        r.meaning, r.props, fmt_pairs(r.examples), fmt_set(r.patterns),
        fmt_set(r.rules_used), fmt_set(r.rules_add), fmt_set(r.rules_del),
        fmt_set(r.rules_addatom), fmt_set(r.rules_delatom),
        fmt_set(r.ifc_mention), fmt_set(r.ifc_update),
        fmt_set(r.ifc_create), fmt_set(r.ifc_remove),
        fmt_set(r.roles["C"]), fmt_set(r.roles["R"]),
        fmt_set(r.roles["U"]), fmt_set(r.roles["D"]),
        "",
    ]


def row_of_rul(u):
    return [
        u.name, u.context, fmt_set(u.origin), u.purpose, u.meaning,
        u.term, u.enforcement, fmt_set(u.maintainers),
        " | ".join(sorted(u.messages)), u.violation, "",
        fmt_set(u.patterns),
        fmt_set(u.rels_used), fmt_set(u.rels_written), fmt_set(u.concepts_touched),
    ]


def fresh_rows(cons, rels, ruls):
    return {
        "Concepts": {c.key: row_of_con(c) for c in cons.values()},
        "Relations": {r.key: row_of_rel(r) for r in rels.values()},
        "Rules": {u.key: row_of_rul(u) for u in ruls.values()},
    }


def label_of(sheet, key):
    if sheet == "Relations":
        return f"{key[0]}[{key[1]}*{key[2]}]"
    return key[0]


# ---------------------------------------------------------------- xlsx I/O

def read_sheet(path, sheet):
    """Read one sheet into {key: row-values-in-canonical-order}. Columns are
    located by their header text, so older column orders still read. Returns {}
    when the file or the sheet does not exist."""
    from openpyxl import load_workbook
    spec = SHEETS[sheet]
    if not Path(path).exists():
        return {}
    wb = load_workbook(path, read_only=True)
    if sheet not in wb.sheetnames:
        return {}
    ws = wb[sheet]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        head = [(str(v).strip() if v is not None else "") for v in next(rows_iter)]
    except StopIteration:
        return {}
    where = [head.index(h) if h in head else None for h in spec["headers"]]
    rows = {}
    for row in rows_iter:
        if not row or not row[0]:
            continue
        vals = []
        for i in where:
            v = row[i] if i is not None and i < len(row) else None
            vals.append(str(v).strip() if v is not None else "")
        rows[tuple(vals[i] for i in spec["key"])] = vals
    return rows


def write_workbook(path, fresh, keep_by_sheet, context):
    """Write the whole workbook: three data sheets plus the Legend. Rows of
    contexts other than `context` pass through untouched; manual columns
    survive when the scan has no value; Origin/Resources fill gaps only."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in ("Concepts", "Relations", "Rules"):
        spec = SHEETS[sheet]
        ws = wb.create_sheet(sheet)
        ws.append(spec["headers"])
        for c in ws[1]:
            c.font = Font(bold=True)
        rows = dict(fresh.get(sheet, {}))
        keep = keep_by_sheet.get(sheet) or {}
        for key, old in keep.items():
            if key in rows:
                for i in spec["manual"]:
                    if not rows[key][i] and old[i]:
                        rows[key][i] = old[i]
                for i in spec["keep"]:
                    if old[i]:
                        rows[key][i] = old[i]
            else:
                rows[key] = old
                if old[spec["context_col"]] == context:
                    print(f"  behouden, niet in scripts ({sheet}): {label_of(sheet, key)}")
        ctx_col = spec["context_col"]
        for key in sorted(rows, key=lambda k: (rows[k][ctx_col].lower(),) +
                          tuple(str(p).lower() for p in k)):
            ws.append(rows[key])
        for i, w in enumerate(spec["widths"], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = spec["freeze"]
        ws.auto_filter.ref = ws.dimensions

    lg = wb.create_sheet("Legend")
    lg.append(["Column / topic", "Explanation"])
    for c in lg[1]:
        c.font = Font(bold=True)
    for item in LEGEND:
        lg.append(list(item))
        if item[1] == "":
            for c in lg[lg.max_row]:
                c.font = Font(bold=True)
    lg.column_dimensions["A"].width = 34
    lg.column_dimensions["B"].width = 110
    for row in lg.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


# ---------------------------------------------------------------- commands

def norm_cell(v):
    return {p.strip() for p in str(v).split(",") if p.strip()}


def read_keep(matrix, seeds):
    """Existing rows per sheet: seed files first, the target workbook wins."""
    keep = {sheet: {} for sheet in SHEETS}
    for src in list(seeds or []) + [matrix]:
        for sheet in SHEETS:
            for key, row in read_sheet(src, sheet).items():
                keep[sheet][key] = row
    return keep


def cmd_scan(args):
    cons, rels, ruls, ctx = parse(args.adl, args.default_crud, args.context)
    fresh = fresh_rows(cons, rels, ruls)
    keep = read_keep(args.matrix, args.seed)
    write_workbook(args.matrix, fresh, keep, ctx)
    print(f"{len(cons)} concepten, {len(rels)} relaties, {len(ruls)} regels "
          f"in CONTEXT {ctx or '(naamloos)'} -> {args.matrix}")


def cmd_check(args):
    cons, rels, ruls, ctx = parse(args.adl, args.default_crud, args.context)
    if not Path(args.matrix).exists():
        print(f"werkboek ontbreekt: {args.matrix}")
        return 1
    fresh = fresh_rows(cons, rels, ruls)
    diffs = 0
    for sheet in ("Concepts", "Relations", "Rules"):
        spec = SHEETS[sheet]
        # Only the rows of the scanned context are compared; other contexts in
        # the workbook have their own scripts (or none, when they are a schema
        # that still has to be built).
        have = {k: v for k, v in read_sheet(args.matrix, sheet).items()
                if v[spec["context_col"]] == ctx}
        new = fresh[sheet]
        for key in sorted(set(new) - set(have)):
            print(f"ONTBREEKT in {sheet}: {label_of(sheet, key)}")
            diffs += 1
        for key in sorted(set(have) - set(new)):
            print(f"NIET in scripts ({sheet}): {label_of(sheet, key)}")
            diffs += 1
        for key in sorted(set(have) & set(new)):
            for i in spec["derived"]:
                a, b = have[key][i], new[key][i]
                if a == "-" and not b:
                    continue
                same = (norm_text(a) == norm_text(b) if i in spec["textcmp"]
                        else norm_cell(a) == norm_cell(b))
                if not same:
                    print(f"VERSCHIL {sheet} {label_of(sheet, key)} / {spec['headers'][i]}:")
                    print(f"  werkboek: {a or '(leeg)'}")
                    print(f"  scripts : {b or '(leeg)'}")
                    diffs += 1
            for i in spec["manual"]:
                a, b = have[key][i], new[key][i]
                if b and a and norm_text(a) != norm_text(b):
                    print(f"VERSCHIL {sheet} {label_of(sheet, key)} / {spec['headers'][i]}:")
                    print(f"  werkboek: {a}")
                    print(f"  scripts : {b}")
                    diffs += 1
    print(f"{diffs} verschil(len) in CONTEXT {ctx}" if diffs
          else f"consistent (CONTEXT {ctx})")
    return 1 if diffs else 0


def cmd_gen(args):
    sheets = {s: read_sheet(args.matrix, s) for s in ("Concepts", "Relations", "Rules")}
    if args.context:
        for s in sheets:
            ctx_col = SHEETS[s]["context_col"]
            sheets[s] = {k: v for k, v in sheets[s].items()
                         if v[ctx_col] == args.context}
    if not any(sheets.values()):
        print(f"niets te genereren{' in CONTEXT ' + args.context if args.context else ''}")
        return 1
    contexts = sorted({v[SHEETS[s]["context_col"]]
                       for s in sheets for v in sheets[s].values()})
    out = []
    for ctx in contexts:
        out.append(f"CONTEXT {ctx or 'Naamloos'} IN DUTCH")
        out.append("")
        for key in sorted((k for k, v in sheets["Concepts"].items()
                           if v[C_CONTEXT] == ctx), key=lambda k: k[0].lower()):
            row = sheets["Concepts"][key]
            name = row[C_NAME]
            if row[C_DEF] and row[C_DEF] != "-":
                out.append(f'CONCEPT {name} "{row[C_DEF]}"')
            else:
                out.append(f'CONCEPT {name} ""  -- definitie ontbreekt nog')
            if row[C_PURPOSE] and row[C_PURPOSE] != "-":
                out.append(f"PURPOSE CONCEPT {name}")
                out.append("{+ " + row[C_PURPOSE] + " +}")
            if row[C_IDENTITY] and row[C_IDENTITY] != "-":
                out.append(f"-- Identity: {row[C_IDENTITY]}")
            if row[C_NOTATION] and row[C_NOTATION] != "-":
                out.append(f"-- Notational conventions: {row[C_NOTATION]}")
            for parent in sorted(norm_cell(row[C_ISA])):
                out.append(f"CLASSIFY {name} ISA {parent}")
            if row[11] and row[11] != "-":
                out.append(f"REPRESENT {name} TYPE {row[11]}")
            out.append("")
        for key in sorted((k for k, v in sheets["Relations"].items()
                           if v[R_CONTEXT] == ctx),
                          key=lambda k: (k[0].lower(), k[1], k[2])):
            row = sheets["Relations"][key]
            name, src, tgt = row[R_NAME], row[R_SOURCE], row[R_TARGET]
            sig = f"{name}[{src}*{tgt}]" if src != tgt else f"{name}[{src}]"
            line = f"RELATION {sig}"
            if row[R_MULT] and row[R_MULT] != "-":
                line += f" [{row[R_MULT]}]"
            out.append(line)
            if row[R_MEANING] and row[R_MEANING] != "-":
                out.append(f'  MEANING "{row[R_MEANING]}"')
            if row[R_PURPOSE] and row[R_PURPOSE] != "-":
                out.append(f"PURPOSE RELATION {name}[{src}*{tgt}]")
                out.append("{+ " + row[R_PURPOSE] + " +}")
            out.append("")
        for key in sorted((k for k, v in sheets["Rules"].items()
                           if v[U_CONTEXT] == ctx), key=lambda k: k[0].lower()):
            row = sheets["Rules"][key]
            name, term = row[U_NAME], row[U_TERM]
            if name.startswith("ENFORCE "):
                if term and term != "-":
                    out.append(f"ENFORCE {term}")
                    out.append("")
                continue
            if not term or term == "-":
                out.append(f"-- RULE {name}: formele term nog te formuleren")
                out.append("")
                continue
            synthetic = name.startswith("RULE ")
            out.append(f"RULE {term}" if synthetic else f'RULE "{name}" : {term}')
            if row[U_MEANING] and row[U_MEANING] != "-":
                out.append(f'MEANING "{row[U_MEANING]}"')
            for msg in (m.strip() for m in row[U_MESSAGE].split("|") if m.strip()):
                if msg != "-":
                    out.append(f'MESSAGE "{msg}"')
            if row[U_VIOLATION] and row[U_VIOLATION] != "-":
                out.append(f"VIOLATION ( {row[U_VIOLATION]} )")
            if not synthetic:
                maint = norm_cell(row[U_MAINT]) - {"-"}
                if maint:
                    roles = ", ".join(f'"{r}"' for r in sorted(maint))
                    out.append(f'ROLE {roles} MAINTAINS "{name}"')
                if row[U_PURPOSE] and row[U_PURPOSE] != "-":
                    out.append(f'PURPOSE RULE "{name}"')
                    out.append("{+ " + row[U_PURPOSE] + " +}")
            out.append("")
        out.append("ENDCONTEXT")
        out.append("")
    Path(args.out).write_text("\n".join(out), encoding="utf-8")
    n = {s: len(sheets[s]) for s in sheets}
    print(f"{n['Concepts']} concepten, {n['Relations']} relaties, {n['Rules']} regels "
          f"uit {len(contexts)} context(en) -> {args.out}")


def read_population(paths):
    """Read Ampersand two-header-row population sheets.

    Row 1 = [SourceConcept] plus relation names (a trailing ~ means flipped),
    row 2 = source concept plus target concepts, row 3+ = data. Returns
    ({(name, src, tgt): [pairs]}, {concept: set of atoms})."""
    from openpyxl import load_workbook

    rel_data, con_data = {}, {}
    for path in paths:
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb:
            rows = ws.iter_rows(values_only=True)
            try:
                head1, head2 = next(rows), next(rows)
            except StopIteration:
                continue
            if not head1 or not head1[0] or not str(head1[0]).startswith("["):
                continue
            src_concept = str(head2[0]).strip()
            cols = []
            for i in range(1, len(head1)):
                name = head1[i]
                if not name or i >= len(head2) or not head2[i]:
                    continue
                name = str(name).strip()
                flip = name.endswith("~")
                name = name.rstrip("~").strip()
                tgt = str(head2[i]).strip()
                key = (name, tgt, src_concept) if flip else (name, src_concept, tgt)
                cols.append((i, key, flip, src_concept if flip else tgt))
            for row in rows:
                if not row or row[0] is None:
                    continue
                a = str(row[0]).strip()
                if a:
                    con_data.setdefault(src_concept, set()).add(a)
                for i, key, flip, concept in cols:
                    if i >= len(row) or row[i] is None:
                        continue
                    b = str(row[i]).strip()
                    if not b:
                        continue
                    rel_data.setdefault(key, []).append((b, a) if flip else (a, b))
                    con_data.setdefault(concept, set()).add(b)
    return rel_data, con_data


def multiplicity_report(pairs, props):
    """Which declared multiplicities does this population contradict?"""
    declared = {p.strip().upper() for p in props.split(",") if p.strip()}
    if "MAP" in declared:
        declared |= {"UNI", "TOT"}
    if "BIJ" in declared:
        declared |= {"INJ", "SUR"}
    srcs, tgts = {}, {}
    for a, b in pairs:
        srcs.setdefault(a, set()).add(b)
        tgts.setdefault(b, set()).add(a)
    broken = []
    if "UNI" in declared:
        bad = [a for a, bs in srcs.items() if len(bs) > 1]
        if bad:
            broken.append(f"UNI geschonden: {len(bad)} bronatomen met meer dan één doel, bv. '{bad[0]}'")
    if "INJ" in declared:
        bad = [b for b, as_ in tgts.items() if len(as_) > 1]
        if bad:
            broken.append(f"INJ geschonden: {len(bad)} doelatomen met meer dan één bron, bv. '{bad[0]}'")
    return broken


def isa_linked(a, b, parents):
    """Is a related to b through the ISA order (either direction, transitively)?"""
    def ancestors(x):
        seen, todo = set(), [x]
        while todo:
            y = todo.pop()
            for p in parents.get(y, ()):
                if p not in seen:
                    seen.add(p)
                    todo.append(p)
        return seen
    return a in ancestors(b) or b in ancestors(a)


def cmd_pop(args):
    if not Path(args.matrix).exists():
        print(f"werkboek ontbreekt: {args.matrix}")
        return 1
    rel_data, con_data = read_population(args.population)
    if not rel_data and not con_data:
        print("geen populatiesheets herkend (verwacht het twee-koprij-formaat)")
        return 1

    from openpyxl import load_workbook
    wb = load_workbook(args.matrix)
    ws_rel, ws_con = wb["Relations"], wb["Concepts"]
    contexts = {str(r[R_CONTEXT].value or "").strip()
                for r in ws_rel.iter_rows(min_row=2) if r[R_NAME].value}
    contexts |= {str(r[C_CONTEXT].value or "").strip()
                 for r in ws_con.iter_rows(min_row=2) if r[C_NAME].value}
    ctx = args.context or (contexts.pop() if len(contexts) == 1 else None)
    if ctx is None:
        print("het werkboek bevat meer dan één context; kies er een met --context "
              f"({', '.join(sorted(c for c in contexts if c))})")
        return 1

    rel_index = {}
    for row in ws_rel.iter_rows(min_row=2):
        if row[R_NAME].value and str(row[R_CONTEXT].value or "").strip() == ctx:
            rel_index[(str(row[R_NAME].value).strip(),
                       str(row[R_SOURCE].value or "").strip(),
                       str(row[R_TARGET].value or "").strip())] = row
    con_index = {}
    for row in ws_con.iter_rows(min_row=2):
        if row[C_NAME].value and str(row[C_CONTEXT].value or "").strip() == ctx:
            con_index[str(row[C_NAME].value).strip()] = row
    parents = {name: norm_cell(row[C_ISA].value or "") for name, row in con_index.items()}

    gevuld = onbekend = 0
    for key in sorted(rel_data):
        pairs = rel_data[key]
        row = rel_index.get(key)
        sig = f"{key[0]}[{key[1]}*{key[2]}]"
        if row is None:
            print(f"NIET in Relations: {sig} — {len(pairs)} paren in de data")
            onbekend += 1
            continue
        if not row[R_EXAMPLES].value:
            row[R_EXAMPLES].value = fmt_pairs(pairs[:3])
            gevuld += 1
        for msg in multiplicity_report(pairs, str(row[R_MULT].value or "")):
            print(f"MULTIPLICITEIT {sig}: {msg}")
        if args.set_resource:
            row[R_RESOURCES].value = args.set_resource
    for name in sorted(con_data):
        atoms = sorted(con_data[name])
        row = con_index.get(name)
        if row is None:
            print(f"NIET in Concepts: {name} — {len(atoms)} atomen in de data")
            onbekend += 1
            continue
        if not row[C_EXAMPLES].value:
            row[C_EXAMPLES].value = fmt_atoms(atoms[:3])
            gevuld += 1
        if args.set_resource:
            row[C_RESOURCES].value = args.set_resource

    # Atoms shared by two concepts without an ISA link: possible homonym, or a
    # CLASSIFY the model lacks. Only reported, never written.
    names = sorted(con_data)
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            both = con_data[a] & con_data[b]
            if both and not isa_linked(a, b, parents):
                ex = sorted(both)[0]
                print(f"GEDEELDE ATOMEN {a} ∩ {b}: {len(both)} atomen zonder "
                      f"ISA-verband, bv. \"{ex}\"")
    wb.save(args.matrix)
    print(f"{len(rel_data)} relaties en {len(con_data)} concepten in de populatie "
          f"van CONTEXT {ctx}; {gevuld} voorbeeldcellen ingevuld; "
          f"{onbekend} niet in het werkboek")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name, fn in (("scan", cmd_scan), ("check", cmd_check)):
        sp = sub.add_parser(name)
        sp.add_argument("matrix")
        sp.add_argument("adl", nargs="+")
        sp.add_argument("--default-crud", default="cRud")
        sp.add_argument("--context", default=None)
        if name == "scan":
            sp.add_argument("--seed", action="append", default=[],
                            help="ouder werkboek (bv. het vroegere ConceptMatrix.xlsx "
                                 "of RelationMatrix.xlsx) waarvan de rijen als "
                                 "basis dienen; herhaalbaar")
        sp.set_defaults(fn=fn)
    sp = sub.add_parser("gen")
    sp.add_argument("matrix")
    sp.add_argument("out")
    sp.add_argument("--context", default=None,
                    help="beperk de generatie tot deze context")
    sp.set_defaults(fn=cmd_gen)
    sp = sub.add_parser("pop")
    sp.add_argument("matrix")
    sp.add_argument("population", nargs="+")
    sp.add_argument("--context", default=None,
                    help="de context waarvan de populatie is; verplicht zodra "
                         "het werkboek meer dan één context bevat")
    sp.add_argument("--set-resource", default=None,
                    help="tekst voor de kolom Resource(s) van elke geraakte rij")
    sp.set_defaults(fn=cmd_pop)
    args = ap.parse_args()
    sys.exit(args.fn(args) or 0)


if __name__ == "__main__":
    main()
